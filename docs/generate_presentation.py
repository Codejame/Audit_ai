"""Generate PowerPoint solution presentation deck for the AI Engineer case study."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    PRIMARY = RGBColor(14, 116, 144)      # Dark cyan / teal
    DARK_BG = RGBColor(15, 23, 42)        # Slate 900
    TEXT_DARK = RGBColor(30, 41, 59)      # Slate 800
    TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
    WHITE = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(248, 250, 252)     # Slate 50
    ACCENT = RGBColor(16, 185, 129)       # Emerald 500

    def add_header(slide, title_text, subtitle_text):
        # Header box
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.5))
    tf1 = bg_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "Intelligent Document Extraction & Validation Platform"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    p2 = tf1.add_paragraph()
    p2.text = "AI Engineer Internship | Technical Case Study Submission"
    p2.font.size = Pt(20)
    p2.font.color.rgb = TEXT_MUTED

    p3 = tf1.add_paragraph()
    p3.text = "\nCandidate: Partha Protim Mondal\nTech Stack: FastAPI, PyMuPDF, Vision LLMs (Gemini/OpenAI), SQLite, Bootstrap 5\nLive Deployed Architecture & Comprehensive Test Suite"
    p3.font.size = Pt(15)
    p3.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Objectives
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "1. Problem Statement & Core Objectives", "Addressing manual bottlenecks in financial services document processing")

    box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.8))
    tf2 = box2.text_frame
    tf2.word_wrap = True

    points2 = [
        ("The Challenge: ", "Financial teams receive diverse invoices and financial statements across native and scanned PDFs, JPGs, and PNGs with varying formats, image noise, and complex line-item structures."),
        ("Multi-Document Scope: ", "Platform natively ingests 4 document categories: Invoices, Balance Sheets, Profit & Loss Statements, and Cash Flow Statements without manual transcription."),
        ("Strict Input Control: ", "Enforces strict input validation: rejects unsupported types (.txt, .docx), corrupted files, and documents exceeding 3 pages."),
        ("Automated Reconciliation: ", "Executes domain-specific financial arithmetic (Total = Subtotal + Tax - Discount, Assets = Liabilities + Equity, Net Cash Flow Reconciliation)."),
        ("Production Mandate: ", "End-to-end cloud deployment on free-tier platform (Render/Railway), interactive UI dashboard, and REST APIs with Swagger documentation.")
    ]
    for bold_prefix, text in points2:
        p = tf2.add_paragraph()
        run1 = p.add_run()
        run1.text = bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(14)
        run1.font.color.rgb = PRIMARY

        run2 = p.add_run()
        run2.text = text + "\n"
        run2.font.size = Pt(14)
        run2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 3: System Architecture
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "2. End-to-End System Architecture", "Modular microservice design ensuring clean separation of concerns")

    box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(5.0))
    tf3 = box3.text_frame
    tf3.word_wrap = True

    points3 = [
        ("FastAPI Application Gateway: ", "High-performance asynchronous backend exposing REST endpoints (/api/v1/documents, /api/v1/health) and interactive OpenAPI documentation (/docs)."),
        ("Validation Layer (DocumentValidationService): ", "Pre-extraction filter validating MIME types, image integrity, and page counts (<= 3 pages) with structured error codes."),
        ("OCR & Rendering (OCRService): ", "PyMuPDF (fitz) rasterizes document pages into high-DPI images for multimodal vision models and extracts native text."),
        ("AI Extraction Engine (ExtractionService): ", "Multimodal LLM adapter supporting Google Gemini and OpenAI with strict Pydantic JSON schemas, plus intelligent local fallback parser."),
        ("Financial Math Engine (FinancialValidationService): ", "Validates domain equations, parses accounting bracketed negatives '(1,000)' -> -1000, checks tolerances, and tags PASS/FAIL/NOT_APPLICABLE."),
        ("Persistence & Dashboard: ", "SQLite repository storing structured responses, queryable by document name, surfaced in responsive Jinja2/Bootstrap 5 dashboard.")
    ]
    for bold_prefix, text in points3:
        p = tf3.add_paragraph()
        run1 = p.add_run()
        run1.text = bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = PRIMARY

        run2 = p.add_run()
        run2.text = text + "\n"
        run2.font.size = Pt(13)
        run2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 4: Financial Validation Engine & Domain Formulas
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "3. Financial Calculation & Validation Engine", "Rigorous mathematical reconciliation across all four financial statements")

    box4 = slide4.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(5.0))
    tf4 = box4.text_frame
    tf4.word_wrap = True

    points4 = [
        ("Invoice Calculations: ", "• Quantity × Unit Price ≈ Line Total\n• Sum(Line Totals) ≈ Subtotal\n• Subtotal + Tax Amount - Discount ≈ Total Amount (tolerates rounding variance)\n• Cash Paid - Total Amount ≈ Change"),
        ("Balance Sheet Reconciliation: ", "• Total Capital & Liabilities ≈ Total Assets (evaluated independently for current and comparative prior years)\n• Sum(Capital + Reserves + Deposits + Borrowings) ≈ Total Liabilities\n• Sum(Cash + Investments + Advances + Fixed Assets) ≈ Total Assets"),
        ("Profit & Loss Reconciliation: ", "• Interest Earned + Other Income ≈ Total Income\n• Interest Expended + Operating Expenses + Provisions ≈ Total Expenditure\n• Total Income - Total Expenditure ≈ Net Profit before Minority Interest\n• Current Profit + Brought Forward Profit ≈ Total Available for Appropriation"),
        ("Cash Flow Statement: ", "• Operating Cash Flow + Investing Cash Flow + Financing Cash Flow + FX ≈ Net Increase in Cash\n• Opening Cash + Net Increase in Cash + Adjustments ≈ Closing Cash\n• Accounting format parser converts '(1,018,904,990)' into negative floats automatically.")
    ]
    for bold_prefix, text in points4:
        p = tf4.add_paragraph()
        run1 = p.add_run()
        run1.text = bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = PRIMARY

        run2 = p.add_run()
        run2.text = text + "\n"
        run2.font.size = Pt(13)
        run2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 5: REST API & Frontend Dashboard
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "4. REST API & Web Dashboard Interface", "Intuitive user experience paired with robust backend standards")

    box5 = slide5.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(5.0))
    tf5 = box5.text_frame
    tf5.word_wrap = True

    points5 = [
        ("POST /api/v1/documents/process: ", "Multipart upload endpoint taking document file + document_type. Synchronously processes, executes validations, persists in database, and returns standard JSON."),
        ("GET /api/v1/documents/{document_name}: ", "Retrieves the latest structured JSON result for a given document name. Returns 404 with structured error if not found."),
        ("GET /api/v1/documents: ", "Lists all processed records with metadata and pass/fail indicators for dashboard population."),
        ("GET /api/v1/health: ", "Liveness probe returning service health and UTC timestamp."),
        ("Interactive Dashboard Features: ", "• Drag & drop upload with file format and page count feedback\n• Real-time processing status spinner\n• Searchable documents list table\n• Detailed result view: key-value pairs, line-item tables, validation checks with color-coded badges, and 1-click Raw JSON modal viewer.")
    ]
    for bold_prefix, text in points5:
        p = tf5.add_paragraph()
        run1 = p.add_run()
        run1.text = bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = PRIMARY

        run2 = p.add_run()
        run2.text = text + "\n"
        run2.font.size = Pt(13)
        run2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 6: Testing & Quality Assurance
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "5. Automated Testing & Verification", "20/20 automated tests passing across validation, extraction, and APIs")

    box6 = slide6.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(5.0))
    tf6 = box6.text_frame
    tf6.word_wrap = True

    points6 = [
        ("Input Validation Tests (test_validation.py): ", "• Valid 1-page PDF and JPG inputs pass.\n• Rejection of unsupported types (.txt) with UNSUPPORTED_FILE_TYPE.\n• Rejection of 0-byte empty files and corrupted headers with CORRUPTED_OR_EMPTY_FILE.\n• Rejection of documents > 3 pages with PAGE_LIMIT_EXCEEDED."),
        ("Domain Math & Negative Number Tests (test_extraction.py): ", "• Negative bracket parser '(1,018,904,990)' -> -1018904990.0 verified.\n• Invoice, Balance Sheet, P&L, and Cash Flow reconciliation formulas tested.\n• Intentional variance test verified: flags FAIL and outputs exact numerical variance.\n• Missing fields verify returning NOT_APPLICABLE rather than hallucinating values."),
        ("REST API Integration Tests (test_api.py): ", "• Health check returns 200 healthy.\n• POST /process verifies 200 for valid uploads, 400 for invalid types.\n• GET /{name} verifies retrieval of latest result and 404 handling.")
    ]
    for bold_prefix, text in points6:
        p = tf6.add_paragraph()
        run1 = p.add_run()
        run1.text = bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = PRIMARY

        run2 = p.add_run()
        run2.text = text + "\n"
        run2.font.size = Pt(13)
        run2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 7: Production Scalability & Deployment
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "6. Deployment & Production Scalability", "Architected for seamless migration to enterprise cloud infrastructure")

    box7 = slide7.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(5.0))
    tf7 = box7.text_frame
    tf7.word_wrap = True

    points7 = [
        ("Containerized Deployment: ", "Dockerfile with multi-stage build, Tesseract OCR system packages, and healthcheck. Ready for 1-click deployment on Render, Railway, or Koyeb via render.yaml and Procfile."),
        ("Zero-Downtime Asynchronous Processing: ", "For high-volume production, decouple upload from processing using Celery / Redis task queues with webhook / WebSocket notifications."),
        ("Cloud Storage & PostgreSQL: ", "Store uploaded documents in AWS S3 / Google Cloud Storage and replace SQLite with managed PostgreSQL with read-replicas."),
        ("Human-in-the-Loop (HITL) Review: ", "Flag extractions with confidence < 85% or validation variances > tolerance for one-click accountant verification and manual adjustment."),
        ("Security & Compliance: ", "Zero secrets committed to GitHub; environment-variable injection; automated PII redaction on financial statements.")
    ]
    for bold_prefix, text in points7:
        p = tf7.add_paragraph()
        run1 = p.add_run()
        run1.text = bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = PRIMARY

        run2 = p.add_run()
        run2.text = text + "\n"
        run2.font.size = Pt(13)
        run2.font.color.rgb = TEXT_DARK

    prs.save("docs/solution_presentation.pptx")
    print("Presentation successfully saved to docs/solution_presentation.pptx")

if __name__ == "__main__":
    create_presentation()
